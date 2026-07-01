# -*- coding: utf-8 -*-
"""
Build the project presentation (PowerPoint .pptx) for:
  Smart Paper Mill Production Monitoring Dashboard
  Industrial Training Project - Satia Industries Limited

Generates: docs/Smart_Paper_Mill_Presentation.pptx
Run:       python docs/build_presentation.py
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
DIAG = os.path.join(HERE, "diagrams")
SHOTS = os.path.join(HERE, "screenshots")
OUT = os.path.join(HERE, "Smart_Paper_Mill_Presentation.pptx")

# ---- palette ----
PRIMARY = RGBColor(0x25, 0x63, 0xEB)
NAVY    = RGBColor(0x0F, 0x2A, 0x47)
INK     = RGBColor(0x0F, 0x17, 0x2A)
MUTED   = RGBColor(0x64, 0x74, 0x8B)
LIGHT   = RGBColor(0xF1, 0xF5, 0xF9)
PALE    = RGBColor(0xE8, 0xEF, 0xFB)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
GREEN   = RGBColor(0x16, 0xA3, 0x4A)
AMBER   = RGBColor(0xF5, 0x9E, 0x0B)
RED     = RGBColor(0xDC, 0x26, 0x26)
CYAN    = RGBColor(0x38, 0xBD, 0xF8)
VIOLET  = RGBColor(0x7C, 0x3A, 0xED)

FONT = "Calibri"
FONT_H = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

_page = {"n": 0}


# =====================================================================
# low-level helpers
# =====================================================================
def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, l, t, w, h, color, line=None):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


def _style_run(r, size, color, bold, italic, font):
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font


def rich(p, text, size, color, bold=False, italic=False, font=FONT):
    """Add a paragraph's runs, honouring **bold** inline markers."""
    parts = text.split("**")
    for i, seg in enumerate(parts):
        if seg == "":
            continue
        r = p.add_run()
        r.text = seg
        _style_run(r, size, color, bold or (i % 2 == 1), italic, font)


def textbox(s, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tf


def line(tf, text, size, color, bold=False, italic=False, align=PP_ALIGN.LEFT,
         before=0, after=6, font=FONT, first=False, level=0, bullet=False,
         line_spacing=None):
    p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align
    p.level = level
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    if line_spacing:
        p.line_spacing = line_spacing
    rich(p, text, size, color, bold=bold, italic=italic, font=font)
    if bullet:
        _set_bullet(p, color)
    else:
        _no_bullet(p)
    return p


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buChar", "a:buAutoNum", "a:buFont", "a:buNone"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def _set_bullet(p, color):
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buFont"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    pPr.set("indent", "-228600")
    pPr.set("marL", "228600")
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": u"•"}))


def footer(s):
    _page["n"] += 1
    tf = textbox(s, Inches(0.7), Inches(7.06), Inches(9), Inches(0.32))
    line(tf, "Smart Paper Mill Production Monitoring Dashboard  -  Satia Industries Ltd.",
         9, MUTED, first=True, after=0)
    tf2 = textbox(s, Inches(12.0), Inches(7.06), Inches(1.0), Inches(0.32))
    line(tf2, str(_page["n"]), 9, MUTED, align=PP_ALIGN.RIGHT, first=True, after=0)


def header(s, kicker, title):
    rect(s, 0, 0, Inches(0.16), SH, PRIMARY)
    tf = textbox(s, Inches(0.7), Inches(0.42), Inches(12.0), Inches(0.34))
    line(tf, kicker.upper(), 12.5, PRIMARY, bold=True, first=True, after=0)
    tt = textbox(s, Inches(0.7), Inches(0.78), Inches(12.2), Inches(0.95))
    line(tt, title, 30, NAVY, bold=True, font=FONT_H, first=True, after=0)
    rect(s, Inches(0.72), Inches(1.72), Inches(1.5), Inches(0.055), PRIMARY)


def body_box(s, top=2.0, height=4.8, left=0.7, width=11.9):
    return textbox(s, Inches(left), Inches(top), Inches(width), Inches(height))


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def image_fit(s, path, l, t, w, h, border=None):
    try:
        from PIL import Image
        with Image.open(path) as im:
            iw, ih = im.size
        ar = iw / float(ih)
    except Exception:
        ar = 16 / 9.0
    box_ar = w / float(h)
    if ar >= box_ar:
        nw = w
        nh = int(w / ar)
    else:
        nh = h
        nw = int(h * ar)
    nl = int(l + (w - nw) / 2)
    nt = int(t + (h - nh) / 2)
    pic = s.shapes.add_picture(path, Emu(nl), Emu(nt), width=Emu(nw))
    if border is not None:
        pic.line.color.rgb = border
        pic.line.width = Pt(1)
    pic.shadow.inherit = False
    return pic


# =====================================================================
# slide builders
# =====================================================================
def title_slide():
    s = slide()
    rect(s, 0, 0, SW, SH, NAVY)
    rect(s, 0, 0, SW, Inches(0.22), PRIMARY)
    rect(s, 0, Inches(5.46), SW, Inches(0.04), RGBColor(0x1E, 0x3A, 0x5F))
    # eyebrow
    tf = textbox(s, Inches(0.9), Inches(0.95), Inches(11), Inches(0.4))
    line(tf, "INDUSTRIAL TRAINING PROJECT  -  2026", 14, CYAN, bold=True, first=True, after=0)
    # title
    tt = textbox(s, Inches(0.9), Inches(1.55), Inches(11.6), Inches(2.2))
    line(tt, "Smart Paper Mill", 48, WHITE, bold=True, font=FONT_H, first=True, after=2)
    line(tt, "Production Monitoring Dashboard", 40, WHITE, bold=True, font=FONT_H, after=0)
    # rule
    rect(s, Inches(0.95), Inches(3.95), Inches(2.2), Inches(0.06), CYAN)
    # subtitle
    st = textbox(s, Inches(0.95), Inches(4.2), Inches(11.4), Inches(1.1))
    line(st, "A real-time, role-based web dashboard for monitoring paper production",
         18, RGBColor(0xCB, 0xD5, 0xE1), first=True, after=2)
    line(st, "Industrial training carried out at  SATIA INDUSTRIES LIMITED", 16,
         RGBColor(0x9F, 0xB3, 0xC8), after=0)
    # submitted-by block
    sb = textbox(s, Inches(0.95), Inches(5.7), Inches(8.5), Inches(1.6))
    line(sb, "Submitted by", 13, RGBColor(0x9F, 0xB3, 0xC8), italic=True, first=True, after=2)
    line(sb, "[STUDENT NAME]   (UID: [____________])", 20, WHITE, bold=True, after=4)
    line(sb, "B.E. Computer Science & Engineering  -  Chandigarh University",
         14, RGBColor(0xCB, 0xD5, 0xE1), after=0)
    # right-side date chip
    dt = textbox(s, Inches(9.7), Inches(5.95), Inches(2.7), Inches(0.9), anchor=MSO_ANCHOR.MIDDLE)
    line(dt, "[MONTH] [YEAR]", 16, WHITE, bold=True, align=PP_ALIGN.RIGHT, first=True, after=0)
    notes(s, "Good morning/afternoon. I'm presenting my industrial training project, the Smart "
              "Paper Mill Production Monitoring Dashboard, developed in the context of Satia "
              "Industries Limited - a wheat-straw based writing and printing paper manufacturer. "
              "All data shown is fictional but industry-realistic; no confidential company data is used.")


def agenda_slide():
    s = slide()
    header(s, "Outline", "Agenda")
    left = ["1.  Organisation & domain", "2.  Problem statement", "3.  Objectives & scope",
            "4.  Technology stack", "5.  System architecture", "6.  Database design"]
    right = ["7.  Security & user roles", "8.  Modules & key features", "9.  Real-time monitoring",
             "10. Quality, OEE & alerts", "11. Testing & results", "12. Challenges & future scope"]
    tf = body_box(s, top=2.05, width=5.8)
    for i, t in enumerate(left):
        line(tf, t, 18, INK, first=(i == 0), after=12)
    tf2 = textbox(s, Inches(6.9), Inches(2.05), Inches(5.8), Inches(4.6))
    for i, t in enumerate(right):
        line(tf2, t, 18, INK, first=(i == 0), after=12)
    footer(s)
    notes(s, "Here's the flow: I'll start with the company and the problem, cover objectives and "
              "scope, then walk through the technology, architecture and database. After that the "
              "modules and the real-time and analytics features, and finally testing, results and "
              "future scope.")


def content_slide(kicker, title, bullets, note, sub=None):
    s = slide()
    header(s, kicker, title)
    tf = body_box(s, top=2.0)
    for i, b in enumerate(bullets):
        if isinstance(b, tuple):
            txt, lvl = b
        else:
            txt, lvl = b, 0
        line(tf, txt, 18 if lvl == 0 else 15.5,
             INK if lvl == 0 else MUTED,
             first=(i == 0), after=10 if lvl == 0 else 6,
             level=lvl, bullet=True)
    footer(s)
    notes(s, note)
    return s


def table_slide(kicker, title, headers, rows, note, col_w=None, first_col_bold=True, fs=13):
    s = slide()
    header(s, kicker, title)
    ncol = len(headers)
    nrow = len(rows) + 1
    left, top = Inches(0.7), Inches(2.05)
    width = Inches(11.9)
    height = Inches(0.5 + 0.46 * len(rows))
    tbl = s.shapes.add_table(nrow, ncol, left, top, width, height).table
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(cw)
    # header row
    for j, htext in enumerate(headers):
        c = tbl.cell(0, j)
        c.fill.solid(); c.fill.fore_color.rgb = NAVY
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_top = Pt(3); c.margin_bottom = Pt(3)
        c.margin_left = Pt(8); c.margin_right = Pt(8)
        tf = c.text_frame; tf.word_wrap = True
        line(tf, htext, fs + 0.5, WHITE, bold=True, first=True, after=0)
    # body rows
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = tbl.cell(i + 1, j)
            c.fill.solid()
            c.fill.fore_color.rgb = WHITE if i % 2 == 0 else LIGHT
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_top = Pt(2); c.margin_bottom = Pt(2)
            c.margin_left = Pt(8); c.margin_right = Pt(8)
            tf = c.text_frame; tf.word_wrap = True
            bold = (j == 0 and first_col_bold)
            line(tf, val, fs, INK if not bold else NAVY, bold=bold, first=True, after=0)
    footer(s)
    notes(s, note)
    return s


def image_slide(kicker, title, img, caption, note):
    s = slide()
    header(s, kicker, title)
    path = os.path.join(DIAG, img)
    if os.path.exists(path):
        image_fit(s, path, Inches(0.9), Inches(1.95), Inches(11.5), Inches(4.45))
    else:
        tf = body_box(s)
        line(tf, "[diagram: %s]" % img, 14, MUTED, first=True)
    cap = textbox(s, Inches(0.7), Inches(6.55), Inches(11.9), Inches(0.45))
    line(cap, caption, 12.5, MUTED, italic=True, align=PP_ALIGN.CENTER, first=True, after=0)
    footer(s)
    notes(s, note)
    return s


def screenshot_slide(kicker, title, img, caption, note):
    s = slide()
    header(s, kicker, title)
    path = os.path.join(SHOTS, img)
    if os.path.exists(path):
        image_fit(s, path, Inches(0.9), Inches(1.95), Inches(11.5), Inches(4.4),
                  border=RGBColor(0xE2, 0xE8, 0xF0))
    else:
        line(body_box(s), "[screenshot: %s]" % img, 14, MUTED, first=True)
    cap = textbox(s, Inches(0.7), Inches(6.5), Inches(11.9), Inches(0.5))
    line(cap, caption, 13, MUTED, italic=True, align=PP_ALIGN.CENTER, first=True, after=0)
    footer(s)
    notes(s, note)
    return s


def screenshot_two(kicker, title, imgL, capL, imgR, capR, note):
    s = slide()
    header(s, kicker, title)
    for img, cap, x in ((imgL, capL, 0.7), (imgR, capR, 6.85)):
        path = os.path.join(SHOTS, img)
        if os.path.exists(path):
            image_fit(s, path, Inches(x), Inches(2.05), Inches(5.75), Inches(3.85),
                      border=RGBColor(0xE2, 0xE8, 0xF0))
        cb = textbox(s, Inches(x), Inches(6.05), Inches(5.75), Inches(0.7))
        line(cb, cap, 12.5, MUTED, italic=True, align=PP_ALIGN.CENTER, first=True, after=0)
    footer(s)
    notes(s, note)
    return s


def closing_slide():
    s = slide()
    rect(s, 0, 0, SW, SH, NAVY)
    rect(s, 0, 0, SW, Inches(0.22), PRIMARY)
    tf = textbox(s, Inches(0.9), Inches(2.6), Inches(11.5), Inches(2.0), anchor=MSO_ANCHOR.MIDDLE)
    line(tf, "Thank You", 54, WHITE, bold=True, font=FONT_H, first=True, after=6)
    line(tf, "Questions & Discussion", 22, CYAN, after=0)
    sb = textbox(s, Inches(0.95), Inches(5.5), Inches(11.4), Inches(1.0))
    line(sb, "[STUDENT NAME]  -  B.E. CSE, Chandigarh University", 15,
         RGBColor(0xCB, 0xD5, 0xE1), first=True, after=2)
    line(sb, "Industrial Training Project  -  Satia Industries Limited", 13,
         RGBColor(0x9F, 0xB3, 0xC8), after=0)
    notes(s, "That concludes the presentation. Thank you - I'm happy to take any questions, and I "
              "can demonstrate the running application live: logging in as different roles, entering "
              "a production record and watching the dashboard, reports and alerts update in real time.")


# =====================================================================
# BUILD
# =====================================================================
title_slide()
agenda_slide()

content_slide(
    "Context", "Organisation & Domain",
    [
        "**Satia Industries Ltd.** - a leading agro-residue paper manufacturer producing "
        "writing & printing paper primarily from **wheat straw** and wood pulp.",
        "Production flows through several stages: **pulping -> stock preparation -> paper "
        "machine (PM) -> finishing**, run on multiple paper machines.",
        "Key process parameters: output (tonnes), **GSM**, **moisture**, machine speed, "
        "temperature, running hours and efficiency.",
        "Continuous, round-the-clock operation across **three shifts (A / B / C)** makes "
        "timely visibility of production, quality and machine health business-critical.",
        ("Disclaimer: all figures in this project are simulated and industry-realistic - no "
         "confidential Satia data is used.", 1),
    ],
    "Satia is one of India's larger eco-friendly paper makers, using wheat straw instead of "
    "felling trees. Paper-making is a continuous process across several stages and runs 24x7 in "
    "three shifts, so managers need a single, live view of what every machine is doing. I want to "
    "stress that all the data in the system is fictional but realistic.")

content_slide(
    "Motivation", "Problem Statement",
    [
        "Production is recorded **manually** in shift registers, logbooks and disconnected "
        "spreadsheets kept separately for each section.",
        "**No consolidated, real-time view** of output, quality (GSM / moisture), machine "
        "status, downtime or efficiency across the plant.",
        "Out-of-spec quality and excessive downtime are often **detected late**, after the "
        "shift, when corrective action is harder.",
        "**KPIs such as OEE are tedious to compute** by hand, and management reports take "
        "time to compile from scattered sources.",
        "Data is hard to secure, audit or share role-by-role across the organisation.",
    ],
    "The core problem is that today the data lives on paper and in spreadsheets. There's no single "
    "live picture, problems are caught late, and computing KPIs or generating reports is slow and "
    "manual. That's the gap this project addresses.")

content_slide(
    "Goals", "Objectives",
    [
        "Provide a **centralised, role-based web dashboard** for real-time production monitoring.",
        "Capture production output and **quality parameters (GSM, moisture, speed)** plus "
        "**downtime events**, persisted in a database.",
        "Show **live machine status and telemetry** and raise **threshold-driven alerts** "
        "automatically.",
        "Compute **OEE and key KPIs** and let users **export reports (CSV / PDF)**.",
        "Enforce **secure authentication and access control**, and run on standard, low-cost "
        "student hardware and a normal browser.",
    ],
    "The objectives follow directly from the problem: one centralised role-based dashboard; capture "
    "of output, quality and downtime that actually persists; live status and automatic alerts; OEE "
    "and exportable reports; and proper security - all while staying lightweight enough to run on a "
    "normal laptop.")

content_slide(
    "Boundaries", "Scope",
    [
        "**In scope:** a complete web application - responsive front-end, REST API back-end and a "
        "database - with four user roles and nine functional modules.",
        "**In scope:** authentication, production & quality entry, machine monitoring, maintenance, "
        "downtime logging, reports, alerts, settings and an admin panel.",
        "**Simulated sensors:** live machine telemetry is streamed from the server to mimic a "
        "real sensor feed, since plant PLC/DCS hardware is not available to a student project.",
        "**Out of scope:** direct PLC / DCS / IoT sensor integration, ERP/SAP integration, and a "
        "native mobile application - identified instead as future scope.",
    ],
    "To keep the project achievable, the scope is the full software stack with four roles and nine "
    "modules. The one honest simplification is that real plant sensors are simulated by a server "
    "stream - I don't have access to the mill's control hardware. Deeper integrations are listed as "
    "future work.")

table_slide(
    "Tools", "Technology Stack",
    ["Layer", "Technology", "Why"],
    [
        ["Front-end", "HTML5, CSS3, JavaScript, Chart.js", "Lightweight, no build step, runs anywhere"],
        ["Back-end", "Python - Flask (REST API + static)", "Simple, same-origin API & UI"],
        ["Database", "SQLite (MySQL is production target)", "Zero-config, file-based persistence"],
        ["Security", "bcrypt, PyJWT", "Hashed passwords + stateless JWT auth"],
        ["Real-time", "Server-Sent Events (SSE)", "Live push with zero extra dependencies"],
        ["Testing", "pytest (94 automated tests)", "Reproducible backend verification"],
        ["Docs / Deploy", "python-docx, Matplotlib; gunicorn / Render", "Auto-generated docs; free hosting"],
    ],
    "The stack is deliberately free, open-source and lightweight. A vanilla HTML/CSS/JS front-end "
    "with Chart.js, a Flask back-end that serves both the API and the pages same-origin, SQLite for "
    "persistence, bcrypt and JWT for security, and Server-Sent Events for the live feed - chosen "
    "over Socket.IO because it needs no extra dependencies. Everything runs on a normal laptop.",
    col_w=[2.2, 4.7, 5.0], fs=12.5)

image_slide(
    "Design", "System Architecture", "architecture-diagram.png",
    "Three-tier architecture: browser front-end  ->  Flask REST API + SSE  ->  SQLite database (same origin).",
    "Architecturally it's a clean three-tier design. The browser runs the pages and charts; a single "
    "Flask application serves both the static front-end and the REST API on the same origin and also "
    "pushes the live SSE stream; and SQLite holds all persisted data. Same-origin serving keeps "
    "authentication and deployment simple.")

image_slide(
    "Data", "Database Design", "er-diagram.png",
    "Core tables: users, production_records, maintenance_logs, downtime_events, machines, thresholds, activity_logs.",
    "The database centres on production records, with supporting tables for users, machines, "
    "maintenance logs, downtime events, configurable thresholds and an activity/audit log. Production "
    "records carry the quality fields - GSM, moisture and speed - alongside output, grade and shift.")

table_slide(
    "Security", "User Roles & Access Control",
    ["Role", "Representative permissions"],
    [
        ["Administrator", "Full access; manage users, machines & thresholds; view audit log"],
        ["Plant Manager", "View all dashboards & reports; oversee production, OEE and downtime"],
        ["Shift Supervisor", "Log production, downtime and maintenance for their shift"],
        ["Machine Operator", "Enter production & quality readings; view machine status & alerts"],
    ],
    "Access is role-based with four roles. Security rests on three pillars: bcrypt-hashed passwords, "
    "stateless JWT tokens, and role-required checks on every protected API endpoint. Sensitive "
    "actions are written to an audit log. So an operator can enter readings but cannot, for example, "
    "manage users - that's enforced on the server, not just hidden in the UI.",
    col_w=[3.0, 8.9], fs=14)

content_slide(
    "Functionality", "Modules & Key Features",
    [
        "**Dashboard** - KPI cards (incl. average OEE), production-trend & machine-status charts, "
        "live feed.",
        "**Production** - record output and **GSM / moisture / speed**; out-of-spec values flagged.",
        "**Machines** - colour-coded status board (Running / Idle / Stopped / Maintenance) with live "
        "temperature & efficiency.",
        "**Maintenance & Downtime** - log and track jobs by priority/status; capture stoppages with "
        "auto-computed duration.",
        "**Reports, Alerts, Settings & Admin** - filtered reports with CSV/PDF export, threshold "
        "alerts, dark mode, and user/machine/threshold management.",
    ],
    "Functionally there are nine modules. The dashboard is the landing view with KPIs and charts. "
    "Production captures output and the three quality parameters. Machines is a live colour-coded "
    "status board. Maintenance and downtime track jobs and stoppages. And reports, alerts, settings "
    "and the admin panel round it out - including CSV/PDF export, dark mode and full admin control.")

content_slide(
    "Live", "Real-Time Monitoring",
    [
        "The server **pushes live readings** (temperature, efficiency) to every open page using "
        "**Server-Sent Events** over an authenticated stream.",
        "Machine cards and the dashboard's live indicator **update every few seconds** without "
        "reloading the page; a 'Live' timestamp shows freshness.",
        "**KPIs, reports and alerts are API-driven:** records entered through the app immediately "
        "feed the dashboard KPIs, the daily report and the alerts list.",
        "Pages **degrade gracefully** - if the backend is unavailable they fall back to a seeded "
        "demonstration dataset, so the UI never breaks.",
    ],
    "A flagship feature is real-time monitoring. The Flask server streams live machine readings using "
    "Server-Sent Events on an authenticated channel, so machine cards and the dashboard update every "
    "few seconds with no page reload. I also wired the dashboard KPIs, the daily report and the "
    "alerts to read live from the API - so if you enter a record it shows up immediately - with a "
    "graceful fallback to seeded data when offline.")

content_slide(
    "Analytics", "Quality, OEE & Alerts",
    [
        "Every production record carries **quality readings - GSM, moisture and machine speed** - "
        "validated against configurable limits.",
        "**Thresholds:** GSM 45-120, moisture <= 6.5%, temperature <= 88 deg C, efficiency >= 75%, "
        "downtime <= 60 min (all admin-editable).",
        "**OEE = Availability x Performance x Quality** is computed and charted per machine.",
        "**Threshold-driven alerts** are raised automatically for out-of-spec GSM/moisture and "
        "over-limit downtime, with acknowledge & severity filtering.",
    ],
    "On the analytics side, each record stores GSM, moisture and speed, checked against limits that "
    "the admin can edit. OEE - availability times performance times quality - is computed and charted "
    "per machine. And the alerts page derives warnings live by comparing the actual data against "
    "those thresholds, so an out-of-spec sheet or an over-long stoppage raises an alert you can "
    "acknowledge.")

screenshot_slide(
    "Walkthrough", "Dashboard", "dashboard.png",
    "Live production overview: KPI cards, production-trend and machine-status charts, and quick actions.",
    "This is the dashboard - the landing screen. Six KPI cards including average OEE, a production-trend "
    "line chart against target, a machine-status doughnut, quick actions and a recent-production table. "
    "The KPIs and charts read live from the API, so they reflect data entered through the app.")

screenshot_slide(
    "Walkthrough", "Machine Monitoring", "machines.png",
    "Colour-coded status board - Running / Idle / Stopped / Maintenance - with live temperature & efficiency.",
    "The machine-monitoring page is a colour-coded status board. Summary chips at the top, search and "
    "status filters - including the Stopped state - and a card per machine showing temperature, running "
    "hours, capacity and an efficiency bar. Temperature and efficiency update live via the server stream.")

screenshot_slide(
    "Walkthrough", "Reports & Analytics", "reports.png",
    "Period / machine / shift filters, plant OEE, downtime breakdown, and CSV / PDF export.",
    "The reports page turns the raw data into analytics: summary KPIs, a daily-production chart that reads "
    "live from the API, a downtime-by-reason breakdown, weekly and monthly views and OEE by machine - all "
    "filterable by period, machine and shift, and exportable to CSV or PDF.")

screenshot_two(
    "Walkthrough", "Production Entry & Alerts",
    "production.png", "Production entry with GSM / moisture / speed quality capture.",
    "alerts.png", "Threshold-driven alerts with severity filtering and acknowledgement.",
    "On the left, the production-entry screen captures output plus the quality parameters - GSM, moisture "
    "and speed - which are validated against the configured limits. On the right, the alerts page, which "
    "raises threshold-driven warnings for out-of-spec quality and excessive downtime that a supervisor "
    "can filter by severity and acknowledge.")

image_slide(
    "Plan", "Project Timeline", "timeline-gantt.png",
    "Phased delivery: requirements & design  ->  module build  ->  backend & security  ->  testing, docs & deployment.",
    "The work was delivered in phases - starting with requirements and design, then building the "
    "modules, then the real back-end with authentication and persistence, and finally testing, "
    "documentation and deployment. This Gantt chart summarises that schedule.")

table_slide(
    "Verification", "Testing & Quality Assurance",
    ["Test area", "Coverage"],
    [
        ["Authentication & JWT", "Valid/invalid login, token guard, expiry handling"],
        ["Role-based access control", "Allowed vs. forbidden endpoints per role (403/200)"],
        ["Production / Maintenance / Downtime", "Full create-read-update-delete + validation"],
        ["Quality & thresholds", "GSM / moisture / speed fields and limit checks"],
        ["Live stream & audit", "Authenticated SSE; audit-log entries"],
        ["Role workflows", "End-to-end tasks for each of the four roles"],
    ],
    "Quality is backed by 94 automated tests in pytest - 36 in the core API suite and 58 in a "
    "role-workflow suite - each running against an isolated temporary database. They cover "
    "authentication, role-based access, full CRUD on production, maintenance and downtime, the "
    "quality fields and thresholds, the live stream and the audit log. All 94 pass.",
    col_w=[4.3, 7.6], fs=13.5)

content_slide(
    "Outcome", "Results & Outcomes",
    [
        "A **working, end-to-end web application** with real authentication, role-based access and "
        "**database persistence that survives restarts**.",
        "**Live telemetry** plus API-driven KPIs, reports and alerts across the dashboard.",
        "A **canonical, seeded dataset** (machines, production, maintenance, operators) for a "
        "consistent, repeatable demonstration.",
        "**94/94 automated tests passing**, and a complete documentation set - **SRS, project "
        "report and role-based user manuals**.",
        "**Deployment-ready** for free hosting (Render) via gunicorn, with no paid services.",
    ],
    "The outcome is a complete, working application - not a mock-up. Real login and roles, data that "
    "persists across restarts, live telemetry, a consistent demo dataset, 94 passing tests, full "
    "documentation including an SRS, report and user manuals, and it's ready to deploy free on Render.")

content_slide(
    "Reflection", "Challenges & Learnings",
    [
        "**Simulating real-time data without hardware** - solved cleanly with Server-Sent Events "
        "instead of polling or heavyweight sockets.",
        "**Keeping documentation consistent with code** - a single source of truth (canonical "
        "dataset + build scripts) prevented drift.",
        "**Security details** - e.g. correct JWT subject typing and enforcing role checks on the "
        "server, not just the UI.",
        "**Designing for constraints** - the whole stack had to run on a low-spec laptop and a "
        "normal browser, which shaped every technology choice.",
        "Gained hands-on experience across **full-stack development, REST APIs, auth, testing and "
        "deployment**.",
    ],
    "The main challenges were simulating live data without plant hardware - solved with SSE - and "
    "keeping a large set of documents consistent with evolving code, which I handled with a single "
    "canonical dataset and generator scripts. I also learned a lot about practical security and about "
    "designing within real hardware constraints. Overall it was genuine full-stack experience.")

content_slide(
    "Next", "Future Scope",
    [
        "Integrate **real plant sensors via PLC / DCS / IoT gateways** to replace the simulated feed.",
        "Migrate to **MySQL / PostgreSQL** for multi-user, production-scale deployment.",
        "Add **predictive maintenance** using machine-learning on historical machine data.",
        "Deliver a **mobile app** and **email / SMS alert** notifications for supervisors on the move.",
        "Extend to **multi-plant / multi-line** monitoring with consolidated group dashboards.",
    ],
    "Looking ahead, the natural next steps are connecting real sensors through PLC or IoT gateways, "
    "moving to a production database, adding predictive maintenance with machine learning, a mobile "
    "app with push and SMS alerts, and scaling from one mill to multiple plants.")

content_slide(
    "Summary", "Conclusion",
    [
        "The project delivers a **secure, real-time, role-based monitoring dashboard** that turns "
        "manual, scattered records into a **single live view** of the plant.",
        "It demonstrates a **complete, tested and documented full-stack system** built entirely with "
        "free, open-source tools.",
        "It directly addresses the industry need for **timely production, quality and machine-health "
        "visibility** - and provides a clear path to a sensor-connected production system.",
    ],
    "To conclude: the project replaces manual logbooks with a secure, real-time, role-based dashboard "
    "- a complete, tested and documented full-stack system built with free tools, and a solid "
    "foundation for a fully sensor-connected solution. Thank you.")

closing_slide()

prs.save(OUT)
print("Saved:", OUT)
print("Slides:", len(prs.slides._sldIdLst))
