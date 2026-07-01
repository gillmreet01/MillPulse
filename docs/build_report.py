# -*- coding: utf-8 -*-
"""
Builds the Industrial Training Report (.docx), formatted to MATCH the
final project report's template (A4, Times New Roman, same title page,
heading sizes, figure/table captions and roman -> arabic page numbering).
Content is refreshed to reflect the current implemented system.
"""

import os
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
DIAG = os.path.join(HERE, "diagrams")
OUT = os.path.join(HERE, "Industrial_Training_Report.docx")

doc = Document()

# ---------------- Page setup (A4) ----------------
sec0 = doc.sections[0]
sec0.page_width, sec0.page_height = Mm(210), Mm(297)
sec0.top_margin = sec0.bottom_margin = Mm(25)
sec0.left_margin = Mm(32); sec0.right_margin = Mm(25)

# ---------------- Base styles (Times New Roman) ----------------
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"; normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.5; normal.paragraph_format.space_after = Pt(6)

def style_heading(name, size):
    st = doc.styles[name]
    st.font.name = "Times New Roman"; st.font.size = Pt(size); st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(12); st.paragraph_format.space_after = Pt(6)

style_heading("Heading 1", 16)
style_heading("Heading 2", 14)
style_heading("Heading 3", 12)

cap = doc.styles["Caption"]
cap.font.name = "Times New Roman"; cap.font.size = Pt(10); cap.font.bold = True
cap.font.color.rgb = RGBColor(0, 0, 0); cap.font.italic = False

# ---------------- Field helpers ----------------
def _field(run, instr):
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = instr
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = "[update field]"
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    for el in (b, i, s, t, e):
        run._r.append(el)

def _seq(p, label):
    run = p.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = " SEQ %s \\* ARABIC " % label
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    run._r.append(b); run._r.append(i); run._r.append(e)
    run.font.name = "Times New Roman"; run.font.size = Pt(10); run.bold = True

def _page_field(p):
    run = p.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = "PAGE"
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    run._r.append(b); run._r.append(i); run._r.append(e)

def set_pgnum(section, fmt, start=None):
    sectPr = section._sectPr
    for e in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(e)
    pg = OxmlElement("w:pgNumType"); pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    sectPr.append(pg)

# ---------------- Content helpers ----------------
def para(text):
    p = doc.add_paragraph(text); p.alignment = AL.JUSTIFY; return p

def center(text, size=12, bold=False, italic=False, after=4):
    p = doc.add_paragraph(); p.alignment = AL.CENTER
    r = p.add_run(text); r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.name = "Times New Roman"; p.paragraph_format.space_after = Pt(after); return p

def title_line(text, size=16, after=6):
    return center(text, size, bold=True, after=after)

def h1(text): return doc.add_heading(text, level=1)
def h2(text): return doc.add_heading(text, level=2)

def subhead(label, body):
    p = doc.add_paragraph(); r = p.add_run(label + "  "); r.bold = True; r.font.name = "Times New Roman"
    p.add_run(body); p.alignment = AL.JUSTIFY; p.paragraph_format.space_after = Pt(6); return p

def bullets(items):
    for it in items: doc.add_paragraph(it, style="List Bullet")

def numbered(items):
    for it in items: doc.add_paragraph(it, style="List Number")

def page_break(): doc.add_page_break()

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell(cell, text, bold=False, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]; r = p.add_run(str(text))
    r.bold = bold; r.font.size = Pt(size); r.font.name = "Times New Roman"
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0

def ctable(title, headers, rows):
    cp = doc.add_paragraph(style="Caption")
    cp.add_run("Table ").bold = True
    _seq(cp, "Table")
    rr = cp.add_run(": " + title); rr.bold = True; rr.font.size = Pt(10)
    cp.runs[0].font.size = Pt(10)
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.allow_autofit = True
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, bold=True); shade(t.rows[0].cells[i], "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            set_cell(cells[i], v)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def figure(filename, title, width=150):
    path = os.path.join(DIAG, filename)
    if os.path.exists(path):
        doc.add_picture(path, width=Mm(width)); doc.paragraphs[-1].alignment = AL.CENTER
    p = doc.add_paragraph(style="Caption"); p.alignment = AL.CENTER
    p.add_run("Figure ").bold = True
    _seq(p, "Figure")
    rr = p.add_run(": " + title); rr.bold = True; rr.font.size = Pt(10)
    p.runs[0].font.size = Pt(10)

def toc_field(instr):
    p = doc.add_paragraph(); _field(p.add_run(), instr)

# Footer page numbers (no number on title page; roman front matter)
sec0.different_first_page_header_footer = True
fp = sec0.footer.paragraphs[0]; fp.alignment = AL.CENTER
fp.add_run("Page "); _page_field(fp)
set_pgnum(sec0, "lowerRoman")

# ==================================================================
# TITLE PAGE
# ==================================================================
for _ in range(2): doc.add_paragraph()
title_line("SMART PAPER MILL PRODUCTION MONITORING DASHBOARD", 18, after=18)
center("AN INDUSTRIAL TRAINING REPORT", 14, bold=True, after=14)
center("Submitted by", 14, bold=True, italic=True, after=10)
center("[STUDENT NAME]  (UID: [__________])", 16, bold=True, after=18)
center("in partial fulfillment for the award of the degree of", 14, bold=True, italic=True, after=10)
center("BACHELOR OF ENGINEERING", 16, bold=True, after=2)
center("IN", 14, after=2)
center("COMPUTER SCIENCE & ENGINEERING", 16, bold=True, after=18)
center("Industrial Training carried out at", 12, after=2)
center("SATIA INDUSTRIES LIMITED", 14, bold=True, after=18)
center("Chandigarh University", 14, after=2)
center("[MONTH] [YEAR]", 14, after=2)
page_break()

# ==================================================================
# TOC / LoF / LoT
# ==================================================================
title_line("TABLE OF CONTENTS", 16, after=10)
toc_field('TOC \\o "1-2" \\h \\z \\u')
para("(Right-click and choose 'Update Field' to generate the contents and page numbers.)")
page_break()
title_line("LIST OF FIGURES", 16, after=10)
toc_field('TOC \\h \\z \\c "Figure"')
para("(Right-click and 'Update Field' to populate the list of figures.)")
page_break()
title_line("LIST OF TABLES", 16, after=10)
toc_field('TOC \\h \\z \\c "Table"')
para("(Right-click and 'Update Field' to populate the list of tables.)")

# ==================================================================
# BODY — arabic page numbers restarting at 1
# ==================================================================
body = doc.add_section(WD_SECTION.NEW_PAGE)
body.page_width, body.page_height = Mm(210), Mm(297)
body.top_margin = body.bottom_margin = Mm(25)
body.left_margin = Mm(32); body.right_margin = Mm(25)
body.different_first_page_header_footer = False
set_pgnum(body, "decimal", start=1)

# 1. ABSTRACT
h1("1. Abstract")
para("This report presents the design, development and validation of the Smart Paper Mill Production "
     "Monitoring Dashboard, a web-based application developed during industrial training at Satia "
     "Industries Limited, a manufacturer of writing and printing paper from agricultural residue "
     "(wheat straw) and wood pulp. The system replaces fragmented, paper-based production logging with "
     "a single, real-time, role-based digital platform that consolidates production output, machine "
     "status, quality parameters, downtime and Overall Equipment Effectiveness (OEE). It is "
     "implemented as a Python (Flask) backend with an SQLite database, secured using bcrypt password "
     "hashing and JSON Web Token authentication with role-based access control, and a responsive "
     "HTML5/CSS3/JavaScript front-end using Chart.js. Live machine telemetry is delivered to the "
     "browser using Server-Sent Events. Because live plant-floor integration is outside the scope of "
     "an academic assignment, the system operates on realistic, internally consistent simulated data. "
     "The delivered solution includes secure login, an administrative panel for users, machines and "
     "thresholds, production and downtime data entry, automatic OEE computation, threshold-based "
     "alerts, analytical reports with CSV and PDF export, an audit trail and a dark-mode interface. "
     "The backend is validated by an automated suite of ninety-four tests, all passing.")

# 2. OBJECTIVES
h1("2. Objectives")
numbered([
    "Provide a centralized, role-based dashboard for real-time paper-mill production monitoring.",
    "Capture production data, including grammage, moisture and machine speed, and flag out-of-specification readings.",
    "Compute Overall Equipment Effectiveness automatically from availability, performance and quality.",
    "Deliver live machine telemetry to the browser without manual page reloads.",
    "Generate threshold-based alerts and present them in a dedicated panel.",
    "Provide analytical reports with date, shift and machine filters and CSV/PDF export.",
    "Provide an administrative panel for managing users, machines and thresholds, with an audit log.",
    "Build the solution with free, open-source tools and validate it with an automated test suite.",
])

# 3. PROBLEM STATEMENT
h1("3. Problem Statement")
para("Production data at the plant is recorded manually in shift registers, logbooks and disconnected "
     "spreadsheets maintained separately for each section. This gives rise to a lack of real-time "
     "visibility, laborious consolidation of scattered data, late detection of out-of-specification "
     "quality and an inability to compare the performance of different shifts, machines or time "
     "periods. The core problem is the absence of a single, real-time, role-based digital platform "
     "through which production, quality and downtime information can be recorded, monitored and "
     "analysed in a consistent and timely manner.")

# 4. SYSTEM ARCHITECTURE
h1("4. System Architecture")
para("The system follows a three-tier architecture, shown in Figure 1.")
figure("architecture-diagram.png", "Three-tier system architecture")
subhead("4.1 Presentation Layer.",
        "The user interface, implemented in HTML5, CSS3 and JavaScript, with Chart.js for "
        "visualization. A shared sidebar and a client-side authentication guard provide consistent "
        "navigation and access control across pages.")
subhead("4.2 Application Layer.",
        "A Python (Flask) REST API that handles authentication, role-based authorization, data CRUD, "
        "OEE computation and live telemetry, and also serves the static front-end so that the whole "
        "application runs from a single origin.")
subhead("4.3 Data Layer.",
        "An SQLite database storing users, production records, downtime events, maintenance logs, "
        "machines, thresholds and the activity log. A relational schema based on MySQL is provided as "
        "the production target; the entity-relationship design is shown in Figure 2.")
figure("er-diagram.png", "Entity-relationship diagram of the database schema")
para("The technologies used are summarised in Table 1.")
ctable("Technology stack",
       ["Layer", "Technology", "Purpose"],
       [["Front-end", "HTML5, CSS3, JavaScript", "Structure, styling, interactivity"],
        ["Charts", "Chart.js", "KPI and trend visualisation"],
        ["Backend", "Python (Flask)", "REST API and static file serving"],
        ["Database", "SQLite", "Persistent storage"],
        ["Security", "bcrypt, PyJWT", "Password hashing and token authentication"],
        ["Real-time", "Server-Sent Events", "Live telemetry push"],
        ["Testing", "pytest", "Automated backend tests"]])

# 5. MODULES
h1("5. Modules")
para("The application is organised into the following modules, listed in Table 2.")
ctable("Application modules",
       ["Module", "Function"],
       [["Authentication", "JWT login, role-based access, session handling"],
        ["Dashboard", "KPI cards (incl. average OEE), trend and status charts, live efficiency feed"],
        ["Machine Monitoring", "Status board with live temperature and efficiency via SSE"],
        ["Production Entry", "Records output, GSM, moisture and speed; flags out-of-spec readings"],
        ["Downtime Logging", "Logs events with reason and automatically computed duration"],
        ["Maintenance", "Logs and tracks maintenance jobs by priority and status"],
        ["Reports", "Filterable charts, OEE by machine, CSV and PDF export"],
        ["Alerts", "Threshold-driven alert panel with acknowledgement"],
        ["Admin Panel", "Manage users, machines and thresholds; view the audit log"]])

# 6. WORKING
h1("6. Working")
para("On launching the application, the user signs in and a JSON Web Token is issued. Each protected "
     "request carries this token, and every protected route enforces a role check. After login, the "
     "dashboard presents summary indicators, charts and a live efficiency figure that is updated "
     "every few seconds from a Server-Sent Events stream pushed by the backend. Operators record "
     "production data, including grammage, moisture and speed, through validated forms; readings that "
     "fall outside the configured tolerance are flagged. Supervisors log downtime events, for which "
     "the duration is computed automatically. All data is persisted to the SQLite database, so it "
     "survives a server restart. The Reports module aggregates the data into charts that can be "
     "filtered by period, machine and shift and exported as CSV or PDF. Administrators manage users, "
     "machines and thresholds and review the audit trail through a dedicated administrative panel.")

# 7. ADVANTAGES
h1("7. Advantages")
numbered([
    "Consolidates production, quality, downtime and machine information within a single, secure interface.",
    "Presents information visually through charts and colour-coded indicators for rapid interpretation.",
    "Uses real authentication (bcrypt and JWT) with role-based access control and an audit trail.",
    "Delivers live data to the browser without manual page reloads.",
    "Built entirely with free and open-source tools and runs on standard hardware.",
    "Validated by an automated test suite, improving reliability and maintainability.",
])

# 8. LIMITATIONS
h1("8. Limitations")
numbered([
    "The system operates on simulated data and is not yet connected to live sensors or a plant control system.",
    "Persistence uses SQLite; migration to MySQL is required for a multi-user production deployment.",
    "The application has not yet been deployed to a public cloud host.",
    "Predictive and analytical capabilities, such as equipment-failure prediction, are not included.",
])

# 9. FUTURE SCOPE
h1("9. Future Scope")
numbered([
    "Integrate real sensors and PLCs using industrial protocols such as OPC-UA or Modbus.",
    "Add predictive maintenance using statistical or machine-learning techniques.",
    "Migrate the database to MySQL and deploy to a free cloud tier for a live demonstration.",
    "Add a dedicated quality-inspection module and a read-only viewer role.",
    "Provide a mobile application and automated report scheduling by electronic mail.",
])

# 10. CONCLUSION
h1("10. Conclusion")
para("The industrial training project delivered a functional, secure and tested web-based dashboard "
     "for monitoring the production activities of a paper manufacturing plant. The system consolidates "
     "production, quality, downtime, maintenance and machine information within a single role-based "
     "interface, computes Overall Equipment Effectiveness, delivers live telemetry, raises "
     "threshold-based alerts, provides analytical reports with export and includes a complete "
     "administrative panel and audit trail. The principal deviation from the original specification is "
     "the absence of live plant instrumentation, which was deliberately replaced by a realistic "
     "simulator suitable for an academic setting. The modular design and accompanying database schema "
     "provide a clear path towards a full production deployment, and the project met the objectives "
     "established at its commencement.")

doc.save(OUT)
print("Saved:", OUT)
