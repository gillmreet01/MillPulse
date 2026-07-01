# -*- coding: utf-8 -*-
"""
Shared PowerPoint toolkit for the Smart Paper Mill presentations.

Holds the palette, layout helpers and slide builders used by both
build_presentation.py (full deck) and build_presentation_short.py (condensed).

Presenter identity is read from the environment so the same scripts can be
re-run for any student without editing code:
    SMM_NAME  - full student name        (default: "[STUDENT NAME]")
    SMM_UID   - university ID / UID       (default: "[____________]")
    SMM_DATE  - month & year on the cover (default: "[MONTH] [YEAR]")
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
DIAG = os.path.join(HERE, "diagrams")
SHOTS = os.path.join(HERE, "screenshots")
OUT = os.path.join(HERE, "Smart_Paper_Mill_Presentation.pptx")
SHORT_OUT = os.path.join(HERE, "Smart_Paper_Mill_Presentation_Short.pptx")

# ---- presenter identity (overridable via environment) ----
STUDENT = os.environ.get("SMM_NAME", "[STUDENT NAME]")
UID = os.environ.get("SMM_UID", "[____________]")
PDATE = os.environ.get("SMM_DATE", "[MONTH] [YEAR]")

# ---- palette ----
PRIMARY = RGBColor(0x25, 0x63, 0xEB)
NAVY    = RGBColor(0x0F, 0x2A, 0x47)
INK     = RGBColor(0x0F, 0x17, 0x2A)
MUTED   = RGBColor(0x64, 0x74, 0x8B)
LIGHT   = RGBColor(0xF1, 0xF5, 0xF9)
PALE    = RGBColor(0xE8, 0xEF, 0xFB)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
GREEN   = RGBColor(0x16, 0xA3, 0x4A)
AMBER   = RGBColor(0xF5, 0x9E, 0x0B)
RED     = RGBColor(0xDC, 0x26, 0x26)
CYAN    = RGBColor(0x38, 0xBD, 0xF8)
VIOLET  = RGBColor(0x7C, 0x3A, 0xED)
SLATE   = RGBColor(0x9F, 0xB3, 0xC8)
FAINT   = RGBColor(0xCB, 0xD5, 0xE1)

FONT = "Calibri"
FONT_H = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

_page = {"n": 0}


# =====================================================================
# low-level helpers
# =====================================================================
def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, l, t, w, h, color, line=None):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


def _style_run(r, size, color, bold, italic, font):
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font


def rich(p, text, size, color, bold=False, italic=False, font=FONT):
    """Add a paragraph's runs, honouring **bold** inline markers."""
    parts = text.split("**")
    for i, seg in enumerate(parts):
        if seg == "":
            continue
        r = p.add_run()
        r.text = seg
        _style_run(r, size, color, bold or (i % 2 == 1), italic, font)


def textbox(s, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tf


def line(tf, text, size, color, bold=False, italic=False, align=PP_ALIGN.LEFT,
         before=0, after=6, font=FONT, first=False, level=0, bullet=False,
         line_spacing=None):
    p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align
    p.level = level
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    if line_spacing:
        p.line_spacing = line_spacing
    rich(p, text, size, color, bold=bold, italic=italic, font=font)
    if bullet:
        _set_bullet(p, color)
    else:
        _no_bullet(p)
    return p


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buChar", "a:buAutoNum", "a:buFont", "a:buNone"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def _set_bullet(p, color):
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buFont"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    pPr.set("indent", "-228600")
    pPr.set("marL", "228600")
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": u"•"}))


def footer(s):
    _page["n"] += 1
    tf = textbox(s, Inches(0.7), Inches(7.06), Inches(9), Inches(0.32))
    line(tf, "Smart Paper Mill Production Monitoring Dashboard  -  Satia Industries Ltd.",
         9, MUTED, first=True, after=0)
    tf2 = textbox(s, Inches(12.0), Inches(7.06), Inches(1.0), Inches(0.32))
    line(tf2, str(_page["n"]), 9, MUTED, align=PP_ALIGN.RIGHT, first=True, after=0)


def header(s, kicker, title):
    rect(s, 0, 0, Inches(0.16), SH, PRIMARY)
    tf = textbox(s, Inches(0.7), Inches(0.42), Inches(12.0), Inches(0.34))
    line(tf, kicker.upper(), 12.5, PRIMARY, bold=True, first=True, after=0)
    tt = textbox(s, Inches(0.7), Inches(0.78), Inches(12.2), Inches(0.95))
    line(tt, title, 30, NAVY, bold=True, font=FONT_H, first=True, after=0)
    rect(s, Inches(0.72), Inches(1.72), Inches(1.5), Inches(0.055), PRIMARY)


def body_box(s, top=2.0, height=4.8, left=0.7, width=11.9):
    return textbox(s, Inches(left), Inches(top), Inches(width), Inches(height))


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def image_fit(s, path, l, t, w, h, border=None):
    try:
        from PIL import Image
        with Image.open(path) as im:
            iw, ih = im.size
        ar = iw / float(ih)
    except Exception:
        ar = 16 / 9.0
    box_ar = w / float(h)
    if ar >= box_ar:
        nw = w
        nh = int(w / ar)
    else:
        nh = h
        nw = int(h * ar)
    nl = int(l + (w - nw) / 2)
    nt = int(t + (h - nh) / 2)
    pic = s.shapes.add_picture(path, Emu(nl), Emu(nt), width=Emu(nw))
    if border is not None:
        pic.line.color.rgb = border
        pic.line.width = Pt(1)
    pic.shadow.inherit = False
    return pic


# =====================================================================
# slide builders
# =====================================================================
def title_slide():
    s = slide()
    rect(s, 0, 0, SW, SH, NAVY)
    rect(s, 0, 0, SW, Inches(0.22), PRIMARY)
    rect(s, 0, Inches(5.46), SW, Inches(0.04), RGBColor(0x1E, 0x3A, 0x5F))
    tf = textbox(s, Inches(0.9), Inches(0.95), Inches(11), Inches(0.4))
    line(tf, "INDUSTRIAL TRAINING PROJECT", 14, CYAN, bold=True, first=True, after=0)
    tt = textbox(s, Inches(0.9), Inches(1.55), Inches(11.6), Inches(2.2))
    line(tt, "Smart Paper Mill", 48, WHITE, bold=True, font=FONT_H, first=True, after=2)
    line(tt, "Production Monitoring Dashboard", 40, WHITE, bold=True, font=FONT_H, after=0)
    rect(s, Inches(0.95), Inches(3.95), Inches(2.2), Inches(0.06), CYAN)
    st = textbox(s, Inches(0.95), Inches(4.2), Inches(11.4), Inches(1.1))
    line(st, "A real-time, role-based web dashboard for monitoring paper production",
         18, FAINT, first=True, after=2)
    line(st, "Industrial training carried out at  SATIA INDUSTRIES LIMITED", 16, SLATE, after=0)
    sb = textbox(s, Inches(0.95), Inches(5.7), Inches(8.5), Inches(1.6))
    line(sb, "Submitted by", 13, SLATE, italic=True, first=True, after=2)
    line(sb, STUDENT + "   (UID: " + UID + ")", 20, WHITE, bold=True, after=4)
    line(sb, "B.E. Computer Science & Engineering  -  Chandigarh University", 14, FAINT, after=0)
    dt = textbox(s, Inches(9.7), Inches(5.95), Inches(2.7), Inches(0.9), anchor=MSO_ANCHOR.MIDDLE)
    line(dt, PDATE, 16, WHITE, bold=True, align=PP_ALIGN.RIGHT, first=True, after=0)
    notes(s, "Good morning/afternoon. I'm presenting my industrial training project, the Smart "
              "Paper Mill Production Monitoring Dashboard, developed in the context of Satia "
              "Industries Limited - a wheat-straw based writing and printing paper manufacturer. "
              "All data shown is fictional but industry-realistic; no confidential company data is used.")


def agenda_slide():
    s = slide()
    header(s, "Outline", "Agenda")
    left = ["1.  Organisation & domain", "2.  Problem statement", "3.  Objectives & scope",
            "4.  Technology stack", "5.  System architecture", "6.  Database design"]
    right = ["7.  Security & user roles", "8.  Modules & key features", "9.  Real-time monitoring",
             "10. Quality, OEE & alerts", "11. Testing & results", "12. Challenges & future scope"]
    tf = body_box(s, top=2.05, width=5.8)
    for i, t in enumerate(left):
        line(tf, t, 18, INK, first=(i == 0), after=12)
    tf2 = textbox(s, Inches(6.9), Inches(2.05), Inches(5.8), Inches(4.6))
    for i, t in enumerate(right):
        line(tf2, t, 18, INK, first=(i == 0), after=12)
    footer(s)
    notes(s, "Here's the flow: I'll start with the company and the problem, cover objectives and "
              "scope, then walk through the technology, architecture and database. After that the "
              "modules and the real-time and analytics features, and finally testing, results and "
              "future scope.")


def content_slide(kicker, title, bullets, note, sub=None):
    s = slide()
    header(s, kicker, title)
    tf = body_box(s, top=2.0)
    for i, b in enumerate(bullets):
        if isinstance(b, tuple):
            txt, lvl = b
        else:
            txt, lvl = b, 0
        line(tf, txt, 18 if lvl == 0 else 15.5,
             INK if lvl == 0 else MUTED,
             first=(i == 0), after=10 if lvl == 0 else 6,
             level=lvl, bullet=True)
    footer(s)
    notes(s, note)
    return s


def table_slide(kicker, title, headers, rows, note, col_w=None, first_col_bold=True, fs=13):
    s = slide()
    header(s, kicker, title)
    ncol = len(headers)
    nrow = len(rows) + 1
    left, top = Inches(0.7), Inches(2.05)
    width = Inches(11.9)
    height = Inches(0.5 + 0.46 * len(rows))
    tbl = s.shapes.add_table(nrow, ncol, left, top, width, height).table
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(cw)
    for j, htext in enumerate(headers):
        c = tbl.cell(0, j)
        c.fill.solid(); c.fill.fore_color.rgb = NAVY
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_top = Pt(3); c.margin_bottom = Pt(3)
        c.margin_left = Pt(8); c.margin_right = Pt(8)
        tf = c.text_frame; tf.word_wrap = True
        line(tf, htext, fs + 0.5, WHITE, bold=True, first=True, after=0)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = tbl.cell(i + 1, j)
            c.fill.solid()
            c.fill.fore_color.rgb = WHITE if i % 2 == 0 else LIGHT
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_top = Pt(2); c.margin_bottom = Pt(2)
            c.margin_left = Pt(8); c.margin_right = Pt(8)
            tf = c.text_frame; tf.word_wrap = True
            bold = (j == 0 and first_col_bold)
            line(tf, val, fs, INK if not bold else NAVY, bold=bold, first=True, after=0)
    footer(s)
    notes(s, note)
    return s


def image_slide(kicker, title, img, caption, note):
    s = slide()
    header(s, kicker, title)
    path = os.path.join(DIAG, img)
    if os.path.exists(path):
        image_fit(s, path, Inches(0.9), Inches(1.95), Inches(11.5), Inches(4.45))
    else:
        tf = body_box(s)
        line(tf, "[diagram: %s]" % img, 14, MUTED, first=True)
    cap = textbox(s, Inches(0.7), Inches(6.55), Inches(11.9), Inches(0.45))
    line(cap, caption, 12.5, MUTED, italic=True, align=PP_ALIGN.CENTER, first=True, after=0)
    footer(s)
    notes(s, note)
    return s


def screenshot_slide(kicker, title, img, caption, note):
    s = slide()
    header(s, kicker, title)
    path = os.path.join(SHOTS, img)
    if os.path.exists(path):
        image_fit(s, path, Inches(0.9), Inches(1.95), Inches(11.5), Inches(4.4),
                  border=RGBColor(0xE2, 0xE8, 0xF0))
    else:
        line(body_box(s), "[screenshot: %s]" % img, 14, MUTED, first=True)
    cap = textbox(s, Inches(0.7), Inches(6.5), Inches(11.9), Inches(0.5))
    line(cap, caption, 13, MUTED, italic=True, align=PP_ALIGN.CENTER, first=True, after=0)
    footer(s)
    notes(s, note)
    return s


def screenshot_two(kicker, title, imgL, capL, imgR, capR, note):
    s = slide()
    header(s, kicker, title)
    for img, cap, x in ((imgL, capL, 0.7), (imgR, capR, 6.85)):
        path = os.path.join(SHOTS, img)
        if os.path.exists(path):
            image_fit(s, path, Inches(x), Inches(2.05), Inches(5.75), Inches(3.85),
                      border=RGBColor(0xE2, 0xE8, 0xF0))
        cb = textbox(s, Inches(x), Inches(6.05), Inches(5.75), Inches(0.7))
        line(cb, cap, 12.5, MUTED, italic=True, align=PP_ALIGN.CENTER, first=True, after=0)
    footer(s)
    notes(s, note)
    return s


def closing_slide():
    s = slide()
    rect(s, 0, 0, SW, SH, NAVY)
    rect(s, 0, 0, SW, Inches(0.22), PRIMARY)
    tf = textbox(s, Inches(0.9), Inches(2.6), Inches(11.5), Inches(2.0), anchor=MSO_ANCHOR.MIDDLE)
    line(tf, "Thank You", 54, WHITE, bold=True, font=FONT_H, first=True, after=6)
    line(tf, "Questions & Discussion", 22, CYAN, after=0)
    sb = textbox(s, Inches(0.95), Inches(5.5), Inches(11.4), Inches(1.0))
    line(sb, STUDENT + "  -  B.E. CSE, Chandigarh University", 15, FAINT, first=True, after=2)
    line(sb, "Industrial Training Project  -  Satia Industries Limited", 13, SLATE, after=0)
    notes(s, "That concludes the presentation. Thank you - I'm happy to take any questions, and I "
              "can demonstrate the running application live: logging in as different roles, entering "
              "a production record and watching the dashboard, reports and alerts update in real time.")


def save(path):
    prs.save(path)
    print("Saved:", path)
    print("Slides:", len(prs.slides._sldIdLst))
